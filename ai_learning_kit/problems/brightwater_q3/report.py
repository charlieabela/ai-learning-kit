"""Assemble findings into a report (HTML and Word) with text, tables and a chart.

The kit handles layout and formatting. What the report says, in the "notes"
you pass in, is yours.
"""
from __future__ import annotations

import base64
import html
import io
from datetime import date
from pathlib import Path

import pandas as pd


def _chart(findings: pd.DataFrame) -> bytes:
    import matplotlib.pyplot as plt
    counts = (findings.assign(status=findings["compliant"].map({True: "Compliant", False: "Not compliant"}))
                      .groupby(["band", "status"]).size().unstack(fill_value=0)
                      .reindex(columns=["Compliant", "Not compliant"], fill_value=0))
    fig, ax = plt.subplots(figsize=(6, 3))
    counts.plot.bar(ax=ax, color=["#4a6fa5", "#c0504d"], rot=0, width=0.7)
    ax.set_xlabel("approval band")
    ax.set_ylabel("transactions")
    ax.set_title("Transactions above €5,000 by approval band and compliance")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for c in ax.containers:
        ax.bar_label(c, fontsize=8)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    return buf.getvalue()


def _summary(findings: pd.DataFrame, splits: pd.DataFrame | None) -> list[str]:
    n = len(findings)
    bad = findings[~findings["compliant"]]
    lines = [f"Transactions above €5,000 (euro equivalent) reviewed: {n}, with a combined value of "
             f"€{findings['amount_eur'].sum():,.2f}.",
             f"Compliant with the policy: {n - len(bad)}. Not compliant: {len(bad)}, "
             f"with a combined value of €{bad['amount_eur'].sum():,.2f}."]
    if splits is not None and len(splits):
        lines.append(f"Possible split purchases: {len(splits)} group(s) involving "
                     f"{int(splits['n_payments'].sum())} payments below €5,000.")
    return lines


def compliance_report(findings: pd.DataFrame, splits: pd.DataFrame | None = None, notes: str = "",
                      author: str = "", title: str = "Q3 2026 Expenditure Compliance Review",
                      filename: str = "q3_compliance_report", formats=("html", "docx"),
                      show: bool = True) -> dict:
    """Write a compliance report and return the paths of the files created.

    findings : the DataFrame returned by kit.pipeline.compliance_check()
    splits   : optional DataFrame from kit.pipeline.split_purchases()
    notes    : your own findings and interpretation, as plain text; separate
               paragraphs with a blank line
    formats  : "html" (opens in a browser) and/or "docx" (opens in Word)
    """
    png = _chart(findings)
    summary = _summary(findings, splits)
    bad = findings[~findings["compliant"]][["erp_ref", "posting_date", "vendor_name", "amount_eur", "band", "issues"]]
    paragraphs = [" ".join(p.split()) for p in notes.strip().split("\n\n") if p.strip()] or ["(No findings text was provided.)"]
    out = {}

    if "html" in formats:
        def table_html(df):
            return df.to_html(index=False, border=0, classes="t", float_format=lambda v: f"{v:,.2f}")
        parts = [
            f"<h1>{html.escape(title)}</h1>",
            f"<p class='meta'>Brightwater Logistics Ltd · {html.escape(author or 'Analyst')} · {date.today():%d %B %Y}</p>",
            "<h2>1. Summary</h2><ul>" + "".join(f"<li>{html.escape(s)}</li>" for s in summary) + "</ul>",
            "<h2>2. Findings</h2>" + "".join(f"<p>{html.escape(p)}</p>" for p in paragraphs),
            "<h2>3. Compliance by band</h2>",
            f"<img alt='chart' src='data:image/png;base64,{base64.b64encode(png).decode()}'>",
            "<h2>4. Transactions not compliant with the policy</h2>", table_html(bad),
        ]
        if splits is not None and len(splits):
            parts += ["<h2>5. Possible split purchases</h2>",
                      table_html(splits[["group", "vendor_name", "n_payments", "first_date", "last_date",
                                         "combined_eur", "erp_refs"]])]
        css = ("body{font-family:Calibri,Arial,sans-serif;max-width:900px;margin:24px auto;padding:0 16px;color:#262626}"
               "h1{color:#1F3864}h2{color:#2E5597;border-bottom:1px solid #B4C6E7;padding-bottom:2px}"
               ".meta{color:#595959}img{max-width:100%}table.t{border-collapse:collapse;font-size:13px;width:100%}"
               ".t th{background:#DCE6F2;text-align:left}.t th,.t td{padding:4px 6px;border-bottom:1px solid #ddd}")
        doc = f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title><style>{css}</style></head><body>{''.join(parts)}</body></html>"
        p = Path(f"{filename}.html")
        p.write_text(doc, encoding="utf-8")
        out["html"] = p
        if show:
            try:
                from IPython.display import HTML, display
                display(HTML(f"<div style='border:1px solid #ccc;padding:8px'>{''.join(parts)}</div>"))
            except ImportError:
                pass

    if "docx" in formats:
        from docx import Document
        from docx.shared import Inches, Pt
        d = Document()
        d.styles["Normal"].font.name = "Calibri"
        d.styles["Normal"].font.size = Pt(11)
        d.add_heading(title, level=0)
        d.add_paragraph(f"Brightwater Logistics Ltd · {author or 'Analyst'} · {date.today():%d %B %Y}")
        d.add_heading("1. Summary", level=1)
        for s in summary:
            d.add_paragraph(s, style="List Bullet")
        d.add_heading("2. Findings", level=1)
        for para in paragraphs:
            d.add_paragraph(para)
        d.add_heading("3. Compliance by band", level=1)
        d.add_picture(io.BytesIO(png), width=Inches(5.5))

        def add_table(df):
            t = d.add_table(rows=1, cols=len(df.columns))
            t.style = "Light Grid Accent 1"
            for cell, col in zip(t.rows[0].cells, df.columns):
                cell.text = str(col)
            for _, row in df.iterrows():
                for cell, v in zip(t.add_row().cells, row):
                    cell.text = f"{v:,.2f}" if isinstance(v, float) else str(v)

        d.add_heading("4. Transactions not compliant with the policy", level=1)
        add_table(bad)
        if splits is not None and len(splits):
            d.add_heading("5. Possible split purchases", level=1)
            add_table(splits[["group", "vendor_name", "n_payments", "first_date", "last_date", "combined_eur"]])
        p = Path(f"{filename}.docx")
        d.save(p)
        out["docx"] = p

    for fmt, p in out.items():
        print(f"Report saved: {p} ({fmt})")
    return out
