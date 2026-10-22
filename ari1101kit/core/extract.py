"""Extraction from documents that are hard to read directly: Word, Excel, PDF and email text.

Every function returns ordinary Python structures (strings, lists, dictionaries)
or a pandas DataFrame. Deciding how to store and combine them is your job.
"""
from __future__ import annotations

import re
from email.utils import parsedate_to_datetime

import pandas as pd

from .case import path

_NUMBER_WORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5}


def _amount(text: str) -> float:
    return float(text.replace(",", ""))


# --------------------------------------------------------------------------- policy (DOCX)
def _policy_document():
    from docx import Document  # imported here so the kit loads even without python-docx
    return Document(path("policy"))


def policy_text() -> str:
    """Return the full text of the Expenditure Approval Policy, one paragraph per line."""
    return "\n".join(p.text for p in _policy_document().paragraphs if p.text.strip())


def policy_thresholds() -> list[dict]:
    """Return Table 1 of the policy (approval bands) as a list of dictionaries.

    The values are exactly as written in the document, so every value is text,
    including the number of approvers.
    """
    table = _policy_document().tables[0]
    header = [c.text.strip() for c in table.rows[0].cells]
    return [{h: c.text.strip() for h, c in zip(header, row.cells)} for row in table.rows[1:]]


def _role_names() -> list[str]:
    roles = set()
    for row in policy_thresholds():
        roles.update(r.strip() for r in row["Roles"].split(";"))
    return sorted(roles, key=len, reverse=True)


def policy_rules() -> list[dict]:
    """Extract the numbered rules (R1, R2, ...) and exceptions (E1, ...) from the policy.

    Each rule becomes a dictionary with its id, title and full text. Where the
    text states them, the dictionary also holds structured details extracted
    from the sentence:

    - "applies_to_eur": {"above": ..., "up_to": ...}   amount limits in euro
    - "approvers": {"count": ..., "roles": [...]}       who must approve
    - "window_days" / "deadline_days" and their unit     time limits
    - "rate_eur_per_usd"                                currency conversion rate

    Different rules have different details, so the dictionaries do not all
    have the same keys.
    """
    roles = _role_names()
    rules = []
    pattern = re.compile(r"^(?P<id>[RE]\d+) \((?P<title>[^)]+)\)\.\s*(?P<text>.+)$")
    for para in _policy_document().paragraphs:
        m = pattern.match(para.text.strip())
        if not m:
            continue
        text = m["text"]
        rule = {"rule_id": m["id"], "title": m["title"], "text": text}

        limits = {}
        above = re.search(r"above €([\d,]+)", text)
        up_to = re.search(r"up to €([\d,]+)", text)
        if above:
            limits["above"] = _amount(above[1])
        if up_to:
            limits["up_to"] = _amount(up_to[1])
        if limits:
            rule["applies_to_eur"] = limits

        found_roles = [r for r in roles if r in text]
        count = re.search(r"approval by (\w+) people", text)
        if found_roles:
            ordered = sorted(found_roles, key=text.index)
            rule["approvers"] = {"count": _NUMBER_WORDS.get(count[1], len(ordered)) if count else len(ordered),
                                 "roles": ordered}

        window = re.search(r"within (\d+) (calendar|working) days of", text)
        if window:
            rule["deadline_days"] = int(window[1])
            rule["deadline_unit"] = f"{window[2]} days"
        else:
            window = re.search(r"within (\d+) (calendar|working) days", text)
            if window:
                rule["window_days"] = int(window[1])
                rule["window_unit"] = f"{window[2]} days"

        rate = re.search(r"([\d.]+) EUR per USD", text)
        if rate:
            rule["rate_eur_per_usd"] = float(rate[1])
        rules.append(rule)
    return rules


# --------------------------------------------------------------------------- vendors (XLSX)
def vendors() -> pd.DataFrame:
    """Read the Approved Vendor List from the Excel workbook into a tidy DataFrame.

    The sheet starts with a title and a note above the real header row, so a plain
    pd.read_excel() would produce misleading column names. This function finds
    the header row, renames the columns to snake_case and converts the dates.
    """
    raw = pd.read_excel(path("vendors"), sheet_name="Approved vendors", header=None)
    header_row = raw.index[raw.iloc[:, 0].astype(str).str.strip().eq("Vendor ID")][0]
    df = raw.iloc[header_row + 1:].copy()
    df.columns = [str(c).strip().lower().replace(" ", "_") for c in raw.iloc[header_row]]
    df = df.dropna(how="all").reset_index(drop=True)
    df["approved_since"] = pd.to_datetime(df["approved_since"]).dt.date
    return df


# --------------------------------------------------------------------------- board report (PDF)
def board_report() -> dict:
    """Extract the Q3 Finance Summary (PDF) into its parts.

    Returns a dictionary with:
    - "text": the full text of the report;
    - "stated_high_value_count": the number of large payments the report claims;
    - "stated_threshold": the wording the report uses for "large";
    - "table": the monthly table as a DataFrame.
    """
    from pypdf import PdfReader
    text = "\n".join(page.extract_text() for page in PdfReader(path("board_report")).pages)
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]

    flat = " ".join(lines)
    claim = re.search(r"Of these, (\d+) payments were of (EUR [\d,]+ or more)", flat)

    rows, months = [], {"July", "August", "September", "Total"}
    for i, ln in enumerate(lines):
        if ln in months and i + 3 < len(lines):
            rows.append({"month": ln, "payments": int(lines[i + 1]),
                         "total_value_eur": _amount(lines[i + 2]),
                         "payments_5000_or_more": int(lines[i + 3])})
    return {"text": text,
            "stated_high_value_count": int(claim[1]) if claim else None,
            "stated_threshold": claim[2] if claim else None,
            "table": pd.DataFrame(rows)}


# --------------------------------------------------------------------------- emails (TXT)
def emails() -> list[dict]:
    """Split the email file into messages and extract simple facts from each.

    Each message becomes a dictionary with "from", "to", "date" (ISO format),
    "subject" and "body", plus two lists found by pattern matching in the subject
    and body: "erp_refs" (transaction references such as TX-2026-01231) and
    "amounts_eur" (amounts written as "EUR 7,850").
    """
    raw = path("emails").read_text(encoding="utf-8")
    messages = []
    for block in re.split(r"^=+\s*$", raw, flags=re.M):
        block = block.strip()
        if not block:
            continue
        head, _, body = block.partition("\n\n")
        fields = dict(re.findall(r"^(From|To|Date|Subject): (.*)$", head, flags=re.M))
        both = fields.get("Subject", "") + "\n" + body
        messages.append({
            "from": fields.get("From"),
            "to": fields.get("To"),
            "date": parsedate_to_datetime(fields["Date"]).isoformat() if "Date" in fields else None,
            "subject": fields.get("Subject"),
            "body": body.strip(),
            "erp_refs": re.findall(r"TX-\d{4}-\d{5}", both),
            "amounts_eur": [_amount(a) for a in re.findall(r"EUR ([\d,]+(?:\.\d+)?)", both)],
        })
    return messages
